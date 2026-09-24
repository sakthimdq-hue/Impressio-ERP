# Copyright (c) 2025, MDQ and contributors
# For license information, please see license.txt

import frappe
from frappe.model.document import Document
from frappe.utils import getdate, today
from datetime import date
from calendar import monthrange
import calendar
from impressio.impressio.api.sales_helper import resolve_sales_order_items
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill


class School(Document):

    def on_trash(self):
        self.delete_related_students()

    def delete_related_students(self):
        if not self.school_code:
            return

        students = frappe.get_all(
            "Students", filters={"school_code": self.school_code}, pluck="name"
        )

        for student in students:
            frappe.delete_doc("Students", student, ignore_permissions=True, force=True)


@frappe.whitelist()
def get_school_code_options():
    # 1. Get unique school_codes from Customer where enrollment is empty
    customer_codes = frappe.get_all(
        "Customer",
        filters={
            "custom_school_code": ["not in", ["", None]],
            "custom_enrollment_number": ["in", ["", None]],
        },
        pluck="custom_school_code",
        distinct=True,
    )

    # 2. Get all school_codes that ALREADY exist in the School DocType
    existing_school_docs = frappe.get_all("School", pluck="school_code")

    # 3. Filter: Keep only codes NOT in the School DocType
    existing_set = set(existing_school_docs)
    new_codes = [code for code in customer_codes if code not in existing_set]

    new_codes.sort()

    # 4. Add the empty string option at the start of the list
    # This allows the user to leave the field blank in the UI
    return ["", "Test Code"] + new_codes


# get matched lead data using School_code.


@frappe.whitelist()
def get_school_lead_data(school_code):
    if not school_code:
        frappe.throw("school_code is required")

    # Get Customer linked to this school code
    customer = frappe.get_all(
        "Customer",
        filters={"custom_school_code": school_code, "custom_enrollment_number": ""},
        fields=["*"],
    )
    if not customer:
        return {}

    # Get Lead linked to Customer
    lead = frappe.get_all(
        "Lead",
        # filters={"company_name": customer[0].customer_name,"status":"completed"},
        filters={"company_name": customer[0].customer_name},
        fields=["name"],
    )
    if not lead:
        return {}

    lead_doc = frappe.get_doc("Lead", lead[0].name)

    response = {
        "select_uniform_details": getattr(lead_doc, "select_uniform_details", 0),
        "select_books_details": getattr(lead_doc, "select_books_details", 0),
        "uniform_table": [],
        "uniform_table2": [],
        "books_table": [],
        "grades_data": []
    }

    # -------------------------------
    # LEAD GENERAL DATA (FOR SCHOOL)
    # -------------------------------
    response["lead_general_data"] = {
        # Common
        "common": {
            "uniform_details_checkbox": getattr(
                lead_doc, "custom_select_uniform_details", 0
            ),
            "books_details_checkbox": getattr(
                lead_doc, "custom_select_books_details", 0
            ),
            "company_name": getattr(lead_doc, "company_name", ""),
        },
        # Address
        "address": {
            "street": getattr(lead_doc, "custom_street", ""),
            "city": getattr(lead_doc, "city", ""),
            "state": getattr(lead_doc, "state", ""),
            "country": getattr(lead_doc, "country", ""),
            "pincode": getattr(lead_doc, "custom_pincode", ""),
            "contact_name": getattr(lead_doc, "custom_full_name", ""),
        },
        # School Coordinator Table
        "school_coordinator": [],
    }

    response["customer"] = customer[0]

    if hasattr(lead_doc, "custom_point_of_contact_info_child_table"):
        for row in lead_doc.custom_point_of_contact_info_child_table:
            response["lead_general_data"]["school_coordinator"].append(
                {
                    "poc_name": getattr(row, "poc_name", ""),
                    "email": getattr(row, "email", ""),
                    "contact_number": getattr(row, "poc_contact_no", ""),
                    "alternate_number": "",
                    "role": getattr(row, "role", ""),
                }
            )

    # -------------------------------
    # UNIFORM DATA
    # -------------------------------
    if hasattr(lead_doc, "custom_lead_uniform_child_table"):
        for idx, uniform_row in enumerate(lead_doc.custom_lead_uniform_child_table):

            uniform_data = {
                # ---------- BASIC ----------
                "grade": getattr(uniform_row, "grade", ""),
                "school_given_grade_name": getattr(
                    uniform_row, "school_given_grade_name", ""
                ),
                "sections": getattr(uniform_row, "sections", ""),
                "school_given_section_name": getattr(
                    uniform_row, "school_given_section_name", ""
                ),
                "house_name": getattr(uniform_row, "house_name", ""),
                "house_colour": getattr(uniform_row, "house_colour", ""),
                # ---------- GENDER / STRENGTH ----------
                "male": getattr(uniform_row, "male", 0),
                "female": getattr(uniform_row, "female", 0),
                "grade_strength": getattr(uniform_row, "grade_strength", 0),
                "house_strength": getattr(uniform_row, "house_strength", 0),
                "total_male_strength": getattr(uniform_row, "total_male_strength", 0),
                "total_female_strength": getattr(
                    uniform_row, "total_female_strength", 0
                ),
                # ---------- UNIFORM TYPE FLAGS ----------
                "uniform_type_regular_uniform": getattr(
                    uniform_row, "uniform_type_regular_uniform", 0
                ),
                "uniform_type_sports_uniform": getattr(
                    uniform_row, "uniform_type_sports_uniform", 0
                ),
                "uniform_type_winter_uniform": getattr(
                    uniform_row, "uniform_type_winter_uniform", 0
                ),
                "uniform_type_accessories": getattr(
                    uniform_row, "uniform_type_accessories", 0
                ),
                "uniform_type_others": getattr(uniform_row, "uniform_type_others", 0),
                # ---------- UNIFORM DETAILS ----------
                "regular_uniform_type": getattr(
                    uniform_row, "regular_uniform_type", ""
                ),
                "uniform_type_name": getattr(uniform_row, "uniform_type_name", ""),
                "uniform_type": getattr(uniform_row, "uniform_type", ""),
                # ---------- PRICING ----------
                "price": getattr(uniform_row, "price", 0.0),
                "selling_price": getattr(uniform_row, "selling_price", 0.0),
                "msl": getattr(uniform_row, "msl", ""),
                # ---------- FABRIC ----------
                "fabric_type": getattr(uniform_row, "fabric_type", ""),
                "sub_fabric_category_type": getattr(
                    uniform_row, "sub_fabric_category_type", ""
                ),
                # ---------- MISC ----------
                "remarks": getattr(uniform_row, "remarks", ""),
                # costing table
                "costing_details": [],
            }

            # Uniform costing
            if hasattr(lead_doc, "custom_uniform_costing_details"):
                for cost_row in lead_doc.custom_uniform_costing_details:
                    if getattr(cost_row, "parent_child_row", None) == getattr(
                        uniform_row, "name", None
                    ):
                        uniform_data["costing_details"].append(
                            {
                                "grade": getattr(cost_row, "grade", ""),
                                "existing_uniform": getattr(
                                    cost_row, "existing_uniform", ""
                                ),
                                "product_name": getattr(cost_row, "product_name", ""),
                                "cost_price": getattr(cost_row, "cost_price", 0.0),
                                "fixed_margin": getattr(cost_row, "fixed_margin", 0.0),
                                "organization_price": getattr(
                                    cost_row, "organization_price", 0.0
                                ),
                                "agreed_price_org": getattr(
                                    cost_row, "agreed_price_org", 0.0
                                ),
                                "organization_margin": getattr(
                                    cost_row, "organization_margin", 0.0
                                ),
                                "organization_mrp": getattr(
                                    cost_row, "organization_mrp", 0.0
                                ),
                                "customer_discount": getattr(
                                    cost_row, "customer_discount", 0.0
                                ),
                                "display_price": getattr(
                                    cost_row, "display_price", 0.0
                                ),
                                "gst_inclusiveexclusive": getattr(
                                    cost_row, "gst_inclusiveexclusive", ""
                                ),
                            }
                        )

            response["uniform_table"].append(uniform_data)

    # -------------------------------
    # BOOKS DATA
    # -------------------------------
    if hasattr(lead_doc, "custom_lead_books_child_table"):
        for idx, books_row in enumerate(lead_doc.custom_lead_books_child_table):
            books_data = {
                # ---------- BASIC ----------
                "grade": getattr(books_row, "grade", ""),
                "school_given_grade_name": getattr(
                    books_row, "school_given_grade_name", ""
                ),
                "sections": getattr(books_row, "sections", ""),
                "school_given_section_name": getattr(
                    books_row, "school_given_section_name", ""
                ),
                # ---------- LANGUAGE ----------
                "language": getattr(books_row, "language", ""),
                "language_strength": getattr(books_row, "language_strength", 0),
                # ---------- GROUP SUBJECT ----------
                "group_subject": getattr(books_row, "group_subject", ""),
                "group_subject_strength": getattr(
                    books_row, "group_subject_strength", 0
                ),
                # ---------- BUNDLE ----------
                "bundle_code": getattr(books_row, "bundle_code", ""),
                "bundle_name": getattr(books_row, "bundle_name", ""),
                "bundle_cost_price": getattr(books_row, "bundle_cost_price", 0.0),
                "selling_price": getattr(books_row, "selling_price", 0.0),
                "bundle_msl": getattr(books_row, "bundle_msl", ""),
                # ---------- MISC ----------
                "remarks": getattr(books_row, "remarks", ""),
                # ---------- STRENGTH / HOUSE ----------
                "total_language_strength": getattr(
                    books_row, "total_language_strength", 0
                ),
                "total_group_subject_strength": getattr(
                    books_row, "total_group_subject_strength", 0
                ),
                "grade_strength": getattr(books_row, "grade_strength", 0),
                "house_name": getattr(books_row, "house_name", ""),
                "house_strength": getattr(books_row, "house_strength", 0),
                # ---------- CHILD TABLE ----------
                "costing_details": [],
            }

            # Books costing
            if hasattr(lead_doc, "custom_books_costing_details"):
                for cost_row in lead_doc.custom_books_costing_details:
                    if getattr(cost_row, "parent_child_row", None) == getattr(
                        books_row, "name", None
                    ):
                        books_data["costing_details"].append(
                            {
                                "grade": getattr(cost_row, "grade", ""),
                                "existing_books": getattr(
                                    cost_row, "existing_books", ""
                                ),
                                "bundle_name": getattr(cost_row, "bundle_name", ""),
                                "sub_bundle_name": getattr(
                                    cost_row, "sub_bundle_name", ""
                                ),
                                "product_name": getattr(cost_row, "product_name", ""),
                                "qty": getattr(cost_row, "qty", 0),
                                "cost_price": getattr(cost_row, "cost_price", 0.0),
                                "fixed_margin": getattr(cost_row, "fixed_margin", 0.0),
                                "organization_price": getattr(
                                    cost_row, "organization_price", 0.0
                                ),
                                "agreed_price_org": getattr(
                                    cost_row, "agreed_price_org", 0.0
                                ),
                                "organization_margin_price": getattr(
                                    cost_row, "organization_margin_price", 0.0
                                ),
                                "organization_mrp": getattr(
                                    cost_row, "organization_mrp", 0.0
                                ),
                                "customer_discount": getattr(
                                    cost_row, "customer_discount", 0.0
                                ),
                                "display_price": getattr(
                                    cost_row, "display_price", 0.0
                                ),
                                "gst_inclusiveexclusive": getattr(
                                    cost_row, "gst_inclusiveexclusive", ""
                                ),
                            }
                        )

            response["books_table"].append(books_data)

    # -------------------------------
    # GRADES DATA (DEDUPED)
    # -------------------------------
    grades_set = set()

    for row in response["uniform_table"]:
        grades_set.add(
            (
                row.get("grade", ""),
                row.get("school_given_grade_name", ""),
                row.get("sections", ""),
            )
        )

    for row in response["books_table"]:
        grades_set.add(
            (
                row.get("grade", ""),
                row.get("school_given_grade_name", ""),
                row.get("sections", ""),
            )
        )

    response["grades_data"] = [
        {"grade": g[0], "school_given_grade_name": g[1], "sections": g[2]}
        for g in grades_set
    ]

    return response


@frappe.whitelist()
def get_lead_data(school_code):
    if not school_code:
        frappe.throw("school_code is required")

    # Get Customer linked to this school code
    customer = frappe.get_all(
        "Customer",
        filters={"custom_school_code": school_code, "custom_enrollment_number": ""},
        fields=["custom_school_code"],
    )
    if not customer:
        return {}

    # Get Lead linked to Customer
    lead = frappe.get_all(
        "Lead", filters={"company_name": customer[0].customer_name}, fields=["name"]
    )
    if not lead:
        return {}

    lead_doc = frappe.get_doc("Lead", lead[0].name)

    # 🔹 Return EVERYTHING related to Lead (including child tables)
    return lead_doc.as_dict()






# For Dashboard data

@frappe.whitelist()
def get_school_dashboard(
    school_code,
    start_date=None,
    end_date=None,
    page=1,
    page_size=10
):


    page = int(page)
    page_size = int(page_size)

    offset = (page - 1) * page_size
    # ---------------------------
    # DEFAULT DATE RANGE
    # ---------------------------

    if not start_date:
        start_date = date.today().replace(month=1, day=1)

    if not end_date:
        end_date = today()

    start_date = getdate(start_date)
    end_date = getdate(end_date)

    data = {}

    # ---------------------------
    # 1️⃣ ALL TIME DATA
    # ---------------------------

    total_students = frappe.db.count(
        "Students",
        {"school_code": school_code}
    )

    # guardians count
    total_guardians = frappe.db.sql("""
        SELECT COUNT(DISTINCT guardian)
        FROM `tabStudent Guardians`
        WHERE parent IN (
            SELECT name FROM `tabStudents`
            WHERE school_code = %(school_code)s
        )
    """, {"school_code": school_code})[0][0] or 0

    # total orders
    total_orders = frappe.db.sql("""
        SELECT COUNT(name)
        FROM `tabSales Order`
        WHERE customer IN (
            SELECT customer
            FROM `tabStudents`
            WHERE school_code = %(school_code)s
        )
        AND docstatus = 1
    """, {"school_code": school_code})[0][0] or 0

    # total revenue
    total_sales = frappe.db.sql("""
        SELECT SUM(grand_total)
        FROM `tabSales Order`
        WHERE customer IN (
            SELECT customer
            FROM `tabStudents`
            WHERE school_code = %(school_code)s
        )
        AND docstatus = 1
    """, {"school_code": school_code})[0][0] or 0

    data["all_time"] = {
        "students": total_students,
        "guardians": total_guardians,
        "orders": total_orders,
        "sales_amount": total_sales or 0
    }

    # ---------------------------
    # 2️⃣ FILTERED DATA
    # ---------------------------

    filtered = {
        "sales_count": 0,
        "sales_amount": 0
    }

    result = frappe.db.sql("""
        SELECT
            COUNT(name) as sales_count,
            SUM(grand_total) as sales_amount
        FROM `tabSales Order`
        WHERE customer IN (
            SELECT customer
            FROM `tabStudents`
            WHERE school_code = %(school_code)s
        )
        AND transaction_date BETWEEN %(start)s AND %(end)s
        AND docstatus = 1
    """, {
        "school_code": school_code,
        "start": start_date,
        "end": end_date
    }, as_dict=True)

    if result:
        filtered = result[0]
        filtered["sales_amount"] = filtered.get("sales_amount") or 0

    data["filtered"] = filtered

    # ---------------------------
    # 3️⃣ CHART DATA
    # ---------------------------

    labels = []
    order_values = []
    amount_values = []

    # ---------------------------
    # CASE 1 : SAME MONTH → DAY WISE
    # ---------------------------

    if start_date.year == end_date.year and start_date.month == end_date.month:

        rows = frappe.db.sql("""
            SELECT
                DAY(transaction_date) as day,
                COUNT(name) as orders,
                SUM(grand_total) as amount
            FROM `tabSales Order`
            WHERE customer IN (
                SELECT customer
                FROM `tabStudents`
                WHERE school_code = %(school_code)s
            )
            AND transaction_date BETWEEN %(start)s AND %(end)s
            AND docstatus = 1
            GROUP BY day
        """, {
            "school_code": school_code,
            "start": start_date,
            "end": end_date
        }, as_dict=True)

        row_map = {r.day: r for r in rows}

        total_days = monthrange(start_date.year, start_date.month)[1]

        for d in range(1, total_days + 1):

            labels.append(str(d))

            r = row_map.get(d)

            order_values.append(r.orders if r else 0)
            amount_values.append(r.amount if r else 0)

    # ---------------------------
    # CASE 2 : SAME YEAR → MONTH WISE
    # ---------------------------

    elif start_date.year == end_date.year:

        rows = frappe.db.sql("""
            SELECT
                MONTH(transaction_date) as month,
                COUNT(name) as orders,
                SUM(grand_total) as amount
            FROM `tabSales Order`
            WHERE customer IN (
                SELECT customer
                FROM `tabStudents`
                WHERE school_code = %(school_code)s
            )
            AND transaction_date BETWEEN %(start)s AND %(end)s
            AND docstatus = 1
            GROUP BY month
        """, {
            "school_code": school_code,
            "start": start_date,
            "end": end_date
        }, as_dict=True)

        row_map = {r.month: r for r in rows}

        for m in range(start_date.month, end_date.month + 1):

            labels.append(calendar.month_abbr[m])   # Jan, Feb, Mar

            r = row_map.get(m)

            order_values.append(r.orders if r else 0)
            amount_values.append(r.amount if r else 0)

    # ---------------------------
    # CASE 3 : MULTIPLE YEARS → YEAR WISE
    # ---------------------------

    else:

        rows = frappe.db.sql("""
            SELECT
                YEAR(transaction_date) as year,
                COUNT(name) as orders,
                SUM(grand_total) as amount
            FROM `tabSales Order`
            WHERE customer IN (
                SELECT customer
                FROM `tabStudents`
                WHERE school_code = %(school_code)s
            )
            AND transaction_date BETWEEN %(start)s AND %(end)s
            AND docstatus = 1
            GROUP BY year
        """, {
            "school_code": school_code,
            "start": start_date,
            "end": end_date
        }, as_dict=True)

        row_map = {r.year: r for r in rows}

        for y in range(start_date.year, end_date.year + 1):

            labels.append(str(y))

            r = row_map.get(y)

            order_values.append(r.orders if r else 0)
            amount_values.append(r.amount if r else 0)

    data["charts"] = {
        "labels": labels,
        "orders": order_values,
        "amounts": amount_values
    }

    data["filter"] = {
        "start_date": start_date,
        "end_date": end_date
    }

    # ---------------------------
    # 4️⃣ SALES ORDER LIST (PAGINATION)
    # ---------------------------

    orders = frappe.db.sql("""
            SELECT
                so.name,
                so.customer,
                so.grand_total,
                so.transaction_date,
                so.status,
                so.creation,
                s.name as student_id,
                s.enrollment_number,
                s.first_name,
                s.last_name,
                s.grade,
                s.section,
                s.joining_date
            FROM `tabSales Order` so
            JOIN `tabStudents` s
                ON s.customer = so.customer
            WHERE s.school_code = %(school_code)s
            AND so.transaction_date BETWEEN %(start)s AND %(end)s
            AND so.docstatus = 1
            ORDER BY so.transaction_date DESC, so.creation DESC
            LIMIT %(page_size)s OFFSET %(offset)s
        """, {
            "school_code": school_code,
            "start": start_date,
            "end": end_date,
            "page_size": page_size,
            "offset": offset
        }, as_dict=True)
    

    for o in orders:
        o["items"] = resolve_sales_order_items(o.name)

    total_rows = frappe.db.sql("""
            SELECT COUNT(name)
            FROM `tabSales Order`
            WHERE customer IN (
                SELECT customer
                FROM `tabStudents`
                WHERE school_code = %(school_code)s
            )
            AND transaction_date BETWEEN %(start)s AND %(end)s
            AND docstatus = 1
        """, {
            "school_code": school_code,
            "start": start_date,
            "end": end_date
        })[0][0]

    data["sales_orders"] = orders
    data["pagination"] = {
        "page": page,
        "page_size": page_size,
        "total_rows": total_rows
    }

    return data





@frappe.whitelist()
def download_school_sales_orders(school_code, start_date=None, end_date=None):

    if not start_date:
        start_date = date.today().replace(month=1, day=1)

    if not end_date:
        end_date = today()

    start_date = getdate(start_date)
    end_date = getdate(end_date)

    orders = frappe.db.sql("""
        SELECT
            so.name,
            so.transaction_date,
            so.customer,
            so.grand_total,
            so.custom_gateway_order_id,
            so.custom_gateway_tracking_id,
            s.name as student_id,
            s.first_name,
            s.last_name,
            s.enrollment_number,
            s.grade,
            s.section
        FROM `tabSales Order` so
        JOIN `tabStudents` s
            ON s.customer = so.customer
        WHERE s.school_code = %(school_code)s
        AND so.transaction_date BETWEEN %(start)s AND %(end)s
        AND so.docstatus = 1
        ORDER BY so.transaction_date DESC
    """, {
        "school_code": school_code,
        "start": start_date,
        "end": end_date
    }, as_dict=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "School Sales Orders"

    headers = [
        "Sales Order",
        "Date",
        "Customer",
        "Student",
        "Grade",
        "Section",
        "Enrollment",
        "Amount",
        "Merchant Id",
        "Transaction Id",
        "Item Code",
        "Qty"
    ]

    # HEADER STYLE
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="E7F3FF", end_color="E7F3FF", fill_type="solid")
    center_align = Alignment(vertical="center")

    ws.append(headers)

    for col in ws[1]:
        col.font = header_font
        col.fill = header_fill
        col.alignment = center_align

    for order in orders:

        student_name = " ".join(
            filter(None, [order.first_name, order.last_name])
        )

        date_str = frappe.format(order.transaction_date, {"fieldtype": "Date"})

        # ORDER HEADER ROW
        ws.append([
            order.name,
            date_str,
            order.customer,
            student_name,
            order.grade,
            order.section,
            order.enrollment_number,
            order.grand_total,
            order.custom_gateway_order_id,
            order.custom_gateway_tracking_id,
            "",
            ""
        ])

        items = resolve_sales_order_items(order.name)

        for item in items:

            ws.append([
                order.name,
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                item["item_code"],
                item["qty"]
            ])

    # COLUMN WIDTHS
    widths = [20, 12, 20, 20, 12, 10, 15, 12, 25, 30, 70, 6]

    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64+i)].width = width

    # ALIGNMENT FOR ALL CELLS
    for row in ws.iter_rows():
        ws.row_dimensions[row[0].row].height = 22
        for cell in row:
            cell.alignment = Alignment(vertical="center")

    # SAVE
    from io import BytesIO

    file_stream = BytesIO()
    wb.save(file_stream)

    start_str = frappe.format(start_date, {"fieldtype": "Date"})
    end_str = frappe.format(end_date, {"fieldtype": "Date"})

    frappe.response["filename"] = f"school_sales_orders_{start_str}_to_{end_str}.xlsx"
    frappe.response["filecontent"] = file_stream.getvalue()
    frappe.response["type"] = "binary"