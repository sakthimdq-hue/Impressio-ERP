# Copyright (c) 2025, MDQ and contributors
# For license information, please see license.txt

import frappe
from frappe import _
from frappe.model.document import Document
from frappe.utils import getdate, today
from datetime import date
from calendar import monthrange
import calendar
from impressio.impressio.api.sales_helper import resolve_sales_order_items
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill


class Students(Document):

    def validate(self):
        self.validate_school_code()
        self.validate_unique_school_enrollment()
        self._enforce_single_preferred("student_billing_addresses")
        self._enforce_single_preferred("student_shipping_addresses")

    # ----------------------------------------------------
    # SCHOOL VALIDATION (BY school_code FIELD)
    # ----------------------------------------------------
    def validate_school_code(self):
        if not self.school_code:
            frappe.throw(_("School Code is mandatory"))

        school_code = str(self.school_code).strip()

        if not frappe.db.exists("School", {"school_code": school_code}):
            # Fallback 1: Check if school_code was passed as school_name
            matched_code = frappe.db.get_value("School", {"school_name": school_code}, "school_code")
            # Fallback 2: Check if school_code was passed as primary key (name)
            if not matched_code and frappe.db.exists("School", school_code):
                matched_code = frappe.db.get_value("School", school_code, "school_code")
            if matched_code:
                school_code = matched_code
            else:
                frappe.throw(
                    _("School with School Code '{0}' does not exist").format(school_code)
                )

        # normalize
        self.school_code = school_code

    # ----------------------------------------------------
    # COMPOSITE UNIQUE VALIDATION
    # ----------------------------------------------------
    def validate_unique_school_enrollment(self):
        if not self.enrollment_number:
            frappe.throw(_("Enrollment Number is mandatory"))

        enrollment_number = str(self.enrollment_number).strip()

        exists = frappe.db.exists(
            "Students",
            {
                "school_code": self.school_code,
                "enrollment_number": enrollment_number,
                "name": ["!=", self.name]  # important for UPDATE
            }
        )

        if exists:
            frappe.throw(
                _("Enrollment Number '{0}' already exists for School '{1}'")
                .format(enrollment_number, self.school_code)
            )

        # normalize
        self.enrollment_number = enrollment_number

    # ----------------------------------------------------
    # PREFERRED ADDRESS ENFORCEMENT
    # ----------------------------------------------------
    def _enforce_single_preferred(self, child_table):
        rows = self.get(child_table) or []

        if not rows:
            return

        preferred_rows = [row for row in rows if row.preferred]

        # Multiple preferred → keep first
        if len(preferred_rows) > 1:
            first = preferred_rows[0]
            for row in rows:
                row.preferred = 1 if row == first else 0
            return

        # None preferred → auto-select first
        if len(preferred_rows) == 0:
            rows[0].preferred = 1
 






@frappe.whitelist()
def get_student_dashboard(
    student,
    start_date=None,
    end_date=None,
    page=1,
    page_size=10
):

    page = int(page)
    page_size = int(page_size)

    offset = (page - 1) * page_size

    student_doc = frappe.get_doc("Students", student)
    customer = student_doc.customer

    # ---------------------------
    # DEFAULT DATE RANGE
    # ---------------------------

    if not start_date:
        start_date = date.today().replace(month=1, day=1)

    if not end_date:
        end_date = today()

    start_date = getdate(start_date)
    end_date = getdate(end_date)

    data = {}

    # ---------------------------
    # 1️⃣ ALL TIME DATA
    # ---------------------------

    # guardian count
    total_guardians = frappe.db.count(
        "Student Guardians",
        {"parent": student}
    )

    # total orders
    total_orders = frappe.db.count(
        "Sales Order",
        {
            "customer": customer,
            "docstatus": 1
        }
    )

    # total revenue
    total_sales = frappe.db.sql("""
        SELECT SUM(grand_total)
        FROM `tabSales Order`
        WHERE customer = %(customer)s
        AND docstatus = 1
    """, {"customer": customer})[0][0] or 0

    data["all_time"] = {
        "guardians": total_guardians,
        "orders": total_orders,
        "sales_amount": total_sales
    }

    # ---------------------------
    # 2️⃣ FILTERED DATA
    # ---------------------------

    result = frappe.db.sql("""
        SELECT
            COUNT(name) as sales_count,
            SUM(grand_total) as sales_amount
        FROM `tabSales Order`
        WHERE customer = %(customer)s
        AND transaction_date BETWEEN %(start)s AND %(end)s
        AND docstatus = 1
    """, {
        "customer": customer,
        "start": start_date,
        "end": end_date
    }, as_dict=True)

    filtered = result[0] if result else {}

    filtered["sales_amount"] = filtered.get("sales_amount") or 0

    data["filtered"] = filtered

    # ---------------------------
    # 3️⃣ CHART DATA
    # ---------------------------

    labels = []
    order_values = []
    amount_values = []

    # SAME MONTH → DAY WISE
    if start_date.year == end_date.year and start_date.month == end_date.month:

        rows = frappe.db.sql("""
            SELECT
                DAY(transaction_date) as day,
                COUNT(name) as orders,
                SUM(grand_total) as amount
            FROM `tabSales Order`
            WHERE customer = %(customer)s
            AND transaction_date BETWEEN %(start)s AND %(end)s
            AND docstatus = 1
            GROUP BY day
        """, {
            "customer": customer,
            "start": start_date,
            "end": end_date
        }, as_dict=True)

        row_map = {r.day: r for r in rows}

        total_days = monthrange(start_date.year, start_date.month)[1]

        for d in range(1, total_days + 1):
            labels.append(str(d))
            r = row_map.get(d)
            order_values.append(r.orders if r else 0)
            amount_values.append(r.amount if r else 0)

    # SAME YEAR → MONTH WISE
    elif start_date.year == end_date.year:

        rows = frappe.db.sql("""
            SELECT
                MONTH(transaction_date) as month,
                COUNT(name) as orders,
                SUM(grand_total) as amount
            FROM `tabSales Order`
            WHERE customer = %(customer)s
            AND transaction_date BETWEEN %(start)s AND %(end)s
            AND docstatus = 1
            GROUP BY month
        """, {
            "customer": customer,
            "start": start_date,
            "end": end_date
        }, as_dict=True)

        row_map = {r.month: r for r in rows}

        for m in range(start_date.month, end_date.month + 1):
            labels.append(calendar.month_abbr[m])
            r = row_map.get(m)
            order_values.append(r.orders if r else 0)
            amount_values.append(r.amount if r else 0)

    # MULTI YEAR → YEAR WISE
    else:

        rows = frappe.db.sql("""
            SELECT
                YEAR(transaction_date) as year,
                COUNT(name) as orders,
                SUM(grand_total) as amount
            FROM `tabSales Order`
            WHERE customer = %(customer)s
            AND transaction_date BETWEEN %(start)s AND %(end)s
            AND docstatus = 1
            GROUP BY year
        """, {
            "customer": customer,
            "start": start_date,
            "end": end_date
        }, as_dict=True)

        row_map = {r.year: r for r in rows}

        for y in range(start_date.year, end_date.year + 1):
            labels.append(str(y))
            r = row_map.get(y)
            order_values.append(r.orders if r else 0)
            amount_values.append(r.amount if r else 0)

    data["charts"] = {
        "labels": labels,
        "orders": order_values,
        "amounts": amount_values
    }

    data["filter"] = {
        "start_date": start_date,
        "end_date": end_date
    }

    # ---------------------------
    # 4️⃣ SALES ORDER LIST
    # ---------------------------

    orders = frappe.db.sql("""
        SELECT
            name,
            grand_total,
            transaction_date,
            status,
            creation
        FROM `tabSales Order`
        WHERE customer = %(customer)s
        AND transaction_date BETWEEN %(start)s AND %(end)s
        AND docstatus = 1
        ORDER BY transaction_date DESC, creation DESC
        LIMIT %(page_size)s OFFSET %(offset)s
    """, {
        "customer": customer,
        "start": start_date,
        "end": end_date,
        "page_size": page_size,
        "offset": offset
    }, as_dict=True)

    for o in orders:
        o["items"] = resolve_sales_order_items(o.name)

    total_rows = frappe.db.count(
        "Sales Order",
        {
            "customer": customer,
            "transaction_date": ["between", [start_date, end_date]],
            "docstatus": 1
        }
    )

    data["sales_orders"] = orders

    data["pagination"] = {
        "page": page,
        "page_size": page_size,
        "total_rows": total_rows
    }

    return data





@frappe.whitelist()
def download_student_sales_orders(student, start_date=None, end_date=None):

    if not start_date:
        start_date = date.today().replace(month=1, day=1)

    if not end_date:
        end_date = today()

    start_date = getdate(start_date)
    end_date = getdate(end_date)

    student_doc = frappe.get_doc("Students", student)

    orders = frappe.db.sql("""
        SELECT
            so.name,
            so.transaction_date,
            so.customer,
            so.grand_total,
            so.custom_gateway_order_id,
            so.custom_gateway_tracking_id,
            s.name as student_id,
            s.first_name,
            s.last_name,
            s.enrollment_number,
            s.grade,
            s.section
        FROM `tabSales Order` so
        JOIN `tabStudents` s
            ON s.customer = so.customer
        WHERE so.transaction_date BETWEEN %(start)s AND %(end)s
        AND so.docstatus = 1
        AND s.customer = %(customer)s
        ORDER BY so.transaction_date DESC
    """, {
        "customer": student_doc.customer,
        "start": start_date,
        "end": end_date
    }, as_dict=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "School Sales Orders"

    headers = [
        "Sales Order",
        "Date",
        "Customer",
        "Student",
        "Grade",
        "Section",
        "Enrollment",
        "Amount",
        "Merchant Id",
        "Transaction Id",
        "Item Code",
        "Qty"
    ]

    # HEADER STYLE
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="E7F3FF", end_color="E7F3FF", fill_type="solid")
    center_align = Alignment(vertical="center")

    ws.append(headers)

    for col in ws[1]:
        col.font = header_font
        col.fill = header_fill
        col.alignment = center_align

    for order in orders:

        student_name = " ".join(
            filter(None, [order.first_name, order.last_name])
        )

        date_str = frappe.format(order.transaction_date, {"fieldtype": "Date"})

        # ORDER HEADER ROW
        ws.append([
            order.name,
            date_str,
            order.customer,
            student_name,
            order.grade,
            order.section,
            order.enrollment_number,
            order.grand_total,
            order.custom_gateway_order_id,
            order.custom_gateway_tracking_id,
            "",
            ""
        ])

        items = resolve_sales_order_items(order.name)

        for item in items:

            ws.append([
                order.name,
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                item["item_code"],
                item["qty"]
            ])

    # COLUMN WIDTHS
    widths = [20, 12, 20, 20, 12, 10, 15, 12, 25, 30, 70, 6]

    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64+i)].width = width

    # ALIGNMENT FOR ALL CELLS
    for row in ws.iter_rows():
        ws.row_dimensions[row[0].row].height = 22
        for cell in row:
            cell.alignment = Alignment(vertical="center")

    # SAVE
    from io import BytesIO

    file_stream = BytesIO()
    wb.save(file_stream)

    start_str = frappe.format(start_date, {"fieldtype": "Date"})
    end_str = frappe.format(end_date, {"fieldtype": "Date"})

    frappe.response["filename"] = f"school_sales_orders_{start_str}_to_{end_str}.xlsx"
    frappe.response["filecontent"] = file_stream.getvalue()
    frappe.response["type"] = "binary"


