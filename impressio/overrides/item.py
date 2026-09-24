import frappe
import json
from frappe import _  
from frappe.utils import cstr, flt, getdate 
from openpyxl import load_workbook
from datetime import datetime
import re
from frappe.utils.pdf import get_pdf
from frappe.utils.file_manager import save_file
from io import BytesIO


@frappe.whitelist()
def create_bundle_master(file_url):
    count = 0
    file_doc = frappe.get_doc("File", {"file_url": file_url})
    file_content = file_doc.get_content()

    wb = load_workbook(filename=BytesIO(file_content), data_only=True)
    sheet = wb.active

    created_bundles = set()
    created_sub_bundles = set()
    created_products = set()

    for idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        bundle = row[1]
        sub_bundle = row[3]
        product = row[17]
        gst_hsn_code = row[29]
        school_name,school_code = frappe.db.get_value('School',{'school_name':row[7]},['name','school_code'])
        grade = row[8]

        # ---------------- Bundle ----------------
        if bundle and bundle not in created_bundles:
            if not frappe.db.exists("Item", bundle):
                bundle_doc = frappe.new_doc("Item")
                grade_match = re.search(r"\d+", grade)
                if not grade_match:
                    frappe.throw(f"Invalid grade format: {grade}")

                grade_number = grade_match.group()   # <-- THIS is the key

                base_code = f"{school_code}7F{grade_number}"
                bundle_doc.item_code = base_code.ljust(14, "$")
                bundle_doc.item_name = bundle
                bundle_doc.gst_hsn_code = '62031100'
                bundle_doc.item_group = "Books"
                bundle_doc.is_stock_item = 1
                bundle_doc.custom_school_name = school_name
                bundle_doc.custom_grade = grade
                bundle_doc.custom_code = '7'
                bundle_doc.save(ignore_permissions=True)

                bom = frappe.new_doc('BOM')
                bom.item = bundle_doc.name
            created_bundles.add(bundle)

        # ---------------- Sub-Bundle ----------------
        sub_bundle_key = f"{bundle}::{sub_bundle}"
        if sub_bundle and sub_bundle_key not in created_sub_bundles:
            if not frappe.db.exists("Item", sub_bundle):
                sub_doc = frappe.new_doc("Item")
                sub_doc.item_code = sub_bundle
                sub_doc.item_name = sub_bundle
                sub_doc.gst_hsn_code = '62031100'
                sub_doc.item_group = "Sub Bundle"
                sub_doc.is_stock_item = 1
                sub_doc.custom_school_name = school_name
                sub_doc.custom_grade = grade
                sub_doc.custom_parent_bundle = bundle  # custom field
                sub_doc.save(ignore_permissions=True)

                bom.append("items", {
                    "item_code": sub_doc.name,
                    "qty": 1
                })
                bom.save()

                sub_bom = frappe.new_doc('BOM')
                sub_bom.item = sub_doc.name

            created_sub_bundles.add(sub_bundle_key)

        # ---------------- Product ----------------
        product_key = f"{sub_bundle}::{product}"
        if product and product_key not in created_products:
            if not frappe.db.exists("Item", product):
                prod_doc = frappe.new_doc("Item")
                prod_doc.item_code = product
                prod_doc.item_name = product
                prod_doc.gst_hsn_code = '62031100'
                prod_doc.item_group = "Books"
                prod_doc.is_stock_item = 1
                prod_doc.custom_parent_sub_bundle = sub_bundle
                prod_doc.custom_school_name = school_name
                prod_doc.custom_grade = grade
                prod_doc.save(ignore_permissions=True)

                
                sub_bom.append("items", {
                    "item_code": prod_doc.name,
                    "qty": row[4]
                })
                sub_bom.save()

            created_products.add(product_key)
        count += 1

    frappe.msgprint("Bundle, Sub-Bundle & Products created successfully ✅")

@frappe.whitelist()
def create_uniform(file_url):
    count = 0
    file_doc = frappe.get_doc("File", {"file_url": file_url})
    file_content = file_doc.get_content()

    wb = load_workbook(filename=BytesIO(file_content), data_only=True)
    sheet = wb.active

    if not frappe.db.exists('Item','Uniform'):
        templ_item = frappe.new_doc('Item')
        templ_item.item_code = 'Uniform'
        templ_item.item_name = 'Uniform'
        templ_item.item_group = 'Uniform'
        templ_item.gst_hsn_code = '62046200'
        templ_item.has_variants = 1
        templ_item.custom_school_name = 'TSUSC-TSUS Chennai'
        templ_item.append("attributes", {
            "attribute": "Size"
        })
        templ_item.append("attributes", {
            "attribute": "Colour"
        })
        templ_item.save()
    
    for idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        school_code = frappe.db.get_value('School','TSUSC-TSUS Chennai','school_code')
        uni_doc = frappe.new_doc("Item")
        val = str(row[17])

        if val.isdigit():
            val = val + "X"
        item_code = f"{school_code}1{row[12]}{row[14]}{val}{row[19]}'A'"
        uni_doc.item_code = item_code
        uni_doc.item_name = row[21]
        uni_doc.gst_hsn_code = '62046200'
        uni_doc.item_group = "Uniform"
        uni_doc.is_stock_item = 1
        uni_doc.custom_school_name = 'TSUSC-TSUS Chennai'
        for i in row[8].split(","):
            uni_doc.append("custom_uniform_grade", {
                "grade": i.strip()
            })
        for j in row[9].split(","):
            uni_doc.append("custom_organization_grade", {
                "grade": j.strip()
            })
        uni_doc.custom_sub_category = row[11]
        uni_doc.custom_sub_sub_category = row[13]
        uni_doc.variant_of = 'Uniform'
        uni_doc.append("attributes", {
            "attribute": "Size",
            "variant_of": "Uniform",
            "attribute_value": row[17]
        })
        uni_doc.append("attributes", {
            "attribute": "Colour",
            "variant_of": "Uniform",
            "attribute_value": row[19]
        })
        uni_doc.custom_msl = row[20]
        uni_doc.description = row[24]
        uni_doc.custom_thumbnail_image = row[23]
        uni_doc.custom_old_sku = row[26]
        uni_doc.custom_weight = row[27]
        uni_doc.custom_length = row[28]
        uni_doc.custom_widthwaist = row[29]
        uni_doc.custom_height = row[30]
        uni_doc.custom_inventre_cost_price = row[31]
        # uni_doc.append("taxes", {
        #     "item_tax_template": frappe.db.get_value(
        #         "Item Tax Template",
        #         {"title": ("like", f"%{row[32]}%"),"company": "Inventre Edu Services Pvt Ltd"},
        #         "name"
        #     )
        # })
        uni_doc.custom_inventre_cost_price_gst = row[34]
        uni_doc.custom_category_fixed_margin = row[35]
        uni_doc.custom_suggested_organization_price = row[36]
        uni_doc.custom_agreed_price_org = row[37]
        uni_doc.custom_organization_margin = row[38]
        uni_doc.custom_organization_mrp = row[39]
        uni_doc.custom_customer_discount = row[40]
        uni_doc.custom_display_price = row[41]
        uni_doc.custom_gst_inclusiveexclusive = row[42]
        uni_doc.custom_opening_qty = row[43]
        uni_doc.custom_measure_image = row[44]
        uni_doc.custom_vendor_name = row[45]
        uni_doc.custom_reordering_tat = row[46]
        uni_doc.custom_minimum_order_quantity = row[47]

        if not frappe.db.exists('Item',item_code):
            uni_doc.save(ignore_permissions=True)
            
    frappe.msgprint("Uniform Item created successfully ✅")
    
    
    
@frappe.whitelist()
def bulk_update_item_status(items, status):
    import json

    if isinstance(items, str):
        items = items.strip().strip("'")  # ← strip outer single quotes first
        try:
            items = json.loads(items)
        except (json.JSONDecodeError, ValueError):
            frappe.throw(_("Invalid items format received."))

    if not isinstance(items, list):
        frappe.throw(_("Items must be a list."))

    item_list = list(dict.fromkeys(
        str(i).strip().strip('"').strip("'")
        for i in items if i
    ))

    if status not in ['Enabled', 'Disabled']:
        frappe.throw(_("Invalid status. Must be 'Enabled' or 'Disabled'"))

    success_count = 0
    failed_count = 0
    failed_items = []

    for item_name in item_list:
        if not item_name:
            continue
        try:
            if not frappe.db.exists("Item", item_name):
                failed_items.append(f"{item_name} (not found)")
                failed_count += 1
                continue

            frappe.db.set_value("Item", item_name, "disabled", 1 if status == "Disabled" else 0)
            success_count += 1

        except Exception as e:
            failed_items.append(f"{item_name} ({str(e)})")
            failed_count += 1

    frappe.db.commit()

    return {
        "success_count": success_count,
        "failed_count": failed_count,
        "failed_items": failed_items
    }