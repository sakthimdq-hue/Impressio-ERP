import frappe
from frappe.model.document import Document
from frappe import _
from frappe.utils import now_datetime
from io import BytesIO
import openpyxl
from frappe.utils.file_manager import save_file

# Shipment categorization thresholds
LARGE_DIMENSION_THRESHOLD = 120  # cm
LARGE_CBM_THRESHOLD = 0.5  # cubic meters
LARGE_WEIGHT_THRESHOLD = 30  # kg


class HandoverToLogistics(Document):
    def validate(self):
        """Validate document before saving"""
        # Validate pincode
        if self.pincode and len(str(self.pincode)) != 6:
            frappe.throw(_("Pincode must be 6 digits"))

        # Validate weight
        if self.weight and self.weight <= 0:
            frappe.throw(_("Weight must be greater than 0"))

        # Auto-calculate dimensions if not set
        if self.weight and not (self.length and self.width and self.height):
            self.set_default_dimensions()

        # Calculate CBM and determine shipment category
        self.calculate_packing_materials_cbm()
        self.determine_shipment_category()
        self.update_child_table_totals()  # NEW: Update child table
        
        if self.payment_type != "Prepaid":
            frappe.throw(_("Only Prepaid shipments are allowed. COD is not supported."))
    
        self.payment_type = "Prepaid"
        self.cod_amount = 0

    def set_default_dimensions(self):
        """Set default dimensions based on weight"""
        try:
            weight = float(self.weight or 1)
            # Simple formula: cube root of volume
            default_dim = round((weight * 5000) ** (1/3), 0)  # Volume in cm³
            default_dim = max(10, min(default_dim, 120))  # Between 10cm and 120cm

            if not self.length:
                self.length = default_dim
            if not self.width:
                self.width = default_dim
            if not self.height:
                self.height = default_dim
        except:
            pass

    def calculate_packing_materials_cbm(self):
        """Calculate total CBM from packing materials table"""
        total_cbm = 0
        max_dimension = 0
        total_weight = 0

        # Calculate from packing materials table if present
        if self.packing_materials:
            for item in self.packing_materials:
                # FIRST: Calculate unit_cbm from dimensions if not set
                length = float(item.pm_length or 0)
                width = float(item.pm_width or 0)
                height = float(item.pm_height or 0)
                qty = float(item.qty or 1)
                weight = float(item.pm_weight or 0)
                
                # Calculate unit CBM from dimensions
                if length > 0 and width > 0 and height > 0:
                    unit_cbm = (length * width * height) / 1000000
                    item.unit_cbm = round(unit_cbm, 6)
                else:
                    unit_cbm = float(item.unit_cbm or 0)
                
                # Calculate item's total CBM
                item.total_cbm = round(unit_cbm * qty, 6)
                total_cbm += item.total_cbm

                # Calculate item's total weight
                item.total_weight = round(weight * qty, 3)
                total_weight += item.total_weight

                # Track maximum dimension
                item_max = max(length, width, height)
                if item_max > max_dimension:
                    max_dimension = item_max

        # If no packing materials, calculate from manual dimensions
        if not self.packing_materials or len(self.packing_materials) == 0:
            length = float(self.length or 0)
            width = float(self.width or 0)
            height = float(self.height or 0)

            if length > 0 and width > 0 and height > 0:
                # Convert cm³ to m³
                total_cbm = (length * width * height) / 1000000
                max_dimension = max(length, width, height)

            total_weight = float(self.weight or 0)

        self.total_cbm = round(total_cbm, 6)
        self.max_dimension = max_dimension

        # Update weight if calculated from packing materials
        if self.packing_materials and total_weight > 0:
            self.weight = round(total_weight, 3)
            
        # Update parent dimensions from packing materials
        self.update_parent_dimensions_from_packing()

    def update_child_table_totals(self):
        """Ensure child table totals are calculated"""
        if self.packing_materials:
            for item in self.packing_materials:
                # Recalculate to ensure consistency
                unit_cbm = float(item.unit_cbm or 0)
                qty = float(item.qty or 1)
                weight = float(item.pm_weight or 0)
                
                item.total_cbm = round(unit_cbm * qty, 6)
                item.total_weight = round(weight * qty, 3)

    def update_parent_dimensions_from_packing(self):
        """Update parent dimension fields from packing materials"""
        if not self.packing_materials or len(self.packing_materials) == 0:
            return
            
        total_weight = 0
        max_length = 0
        max_width = 0
        total_height = 0
        
        for item in self.packing_materials:
            qty = float(item.qty or 1)
            length = float(item.pm_length or 0)
            width = float(item.pm_width or 0)
            height = float(item.pm_height or 0)
            weight = float(item.pm_weight or 0)
            
            total_weight += weight * qty
            if length > max_length:
                max_length = length
            if width > max_width:
                max_width = width
            total_height += height * qty
        
        # Update parent fields
        self.weight = round(total_weight, 3)
        self.length = max_length
        self.width = max_width
        self.height = round(total_height, 2)
        
    def _get_carrier_thresholds(self):
        carrier = self.logistics_partner or ""
        if carrier == "Ekart":
            settings = frappe.get_single("Logistics Settings")
            weight_threshold = float(getattr(settings, "ekart_large_weight_threshold", 4) or 4)
            return {"weight": weight_threshold, "cbm": 0.5, "dimension": 120, "has_large": True}
        else:
            # Amazon, Shiprocket — no Large/Non-Large distinction
            return {"weight": 9999, "cbm": 9999, "dimension": 9999, "has_large": False}


    # REPLACE determine_shipment_category:
    def determine_shipment_category(self):
        thresholds = self._get_carrier_thresholds()
        
        if not thresholds["has_large"]:
            self.shipment_category = "Non-Large"
            return

        is_large = False
        if self.max_dimension and self.max_dimension > thresholds["dimension"]:
            is_large = True
        if self.total_cbm and self.total_cbm > thresholds["cbm"]:
            is_large = True
        if self.weight and float(self.weight) > thresholds["weight"]:
            is_large = True
        if not self.packing_materials or len(self.packing_materials) == 0:
            if (float(self.length or 0) > thresholds["dimension"] or
                float(self.width or 0) > thresholds["dimension"] or
                float(self.height or 0) > thresholds["dimension"]):
                is_large = True

        self.shipment_category = "Large" if is_large else "Non-Large"
    
    @frappe.whitelist()
    def create_ekart_shipment(self):
        """Create Ekart shipment from document"""
        from impressio.material_out.api.ekart.utils import create_shipment_for_doc
        return create_shipment_for_doc(self.name)
    
    @frappe.whitelist()
    def track_ekart_shipment(self):
        """Track Ekart shipment"""
        from impressio.material_out.api.ekart.utils import track_shipment
        return track_shipment(self.name)
    
    @frappe.whitelist()
    def download_ekart_label(self):
        """Download Ekart label"""
        from impressio.material_out.api.ekart.utils import download_label
        return download_label(self.name)
    
    def before_save(self):
        """Update status based on Ekart status"""
        if self.ekart_shipment_status:
            self.update_document_status()

    def after_insert(self):
        """After document is inserted, update linked dispatch"""
        self.update_linked_dispatch_status("IAWB Created")

    def on_update(self):
        """When document is updated, sync status to dispatch"""
        if self.dispatch_order:
            iawb_status = self.get_dispatch_iawb_status()
            if iawb_status:
                self.update_linked_dispatch_status(iawb_status)

    def get_dispatch_iawb_status(self):
        """Map logistics status to dispatch IAWB status"""
        status = self.ekart_shipment_status or self.carrier_status or self.status

        status_mapping = {
            "Created": "IAWB Created",
            "Pending": "IAWB Created",
            "Picked Up": "Shipped",
            "In Transit": "Shipped",
            "Out for Delivery": "Shipped",
            "Delivered": "Delivered",
            "RTO": "Shipped",
            "Exception": "Shipped",
            "Cancelled": "IAWB Pending"
        }

        return status_mapping.get(status)

    def update_linked_dispatch_status(self, iawb_status):
        """Update the linked dispatch order's IAWB status"""
        if not self.dispatch_order:
            return

        from impressio.material_out.doctype.dispatch_of_orders.dispatch_of_orders import update_iawb_status
        update_iawb_status(
            self.dispatch_order,
            iawb_status,
            logistics_doc=self.name,
            tracking_number=self.logistics_tracking_number or self.ekart_tracking_number,
            logistics_partner=self.logistics_partner
        )

    def update_document_status(self):
        """Update document status based on Ekart status"""
        status_mapping = {
            "Delivered": "Delivered",
            "In Transit": "In Transit",
            "Out for Delivery": "In Transit",
            "Picked Up": "Handed Over",
            "RTO": "Returned",
            "Cancelled": "Draft"
        }

        mapped_status = status_mapping.get(self.ekart_shipment_status)
        if mapped_status and self.status != mapped_status:
            self.status = mapped_status

    @frappe.whitelist()
    def verify_box_by_scan(self, barcode_or_qr):
        """Verify a box by scanning its barcode/QR code during logistics handover"""
        import json

        if not barcode_or_qr:
            return {"success": False, "message": "No barcode/QR code provided"}

        # Try to parse as JSON (QR code data)
        try:
            qr_data = json.loads(barcode_or_qr)
            item_code = qr_data.get("item_code")
        except (json.JSONDecodeError, TypeError):
            item_code = barcode_or_qr

        # Check if item is in packing materials
        for row in self.packing_materials or []:
            if row.item_code == item_code:
                return {
                    "success": True,
                    "message": f"Box {item_code} verified for IAWB mapping",
                    "item_code": item_code,
                    "item_name": row.item_name,
                    "dimensions": {
                        "l": row.pm_length,
                        "w": row.pm_width,
                        "h": row.pm_height
                    }
                }

        return {"success": False, "message": f"Box {barcode_or_qr} not found in shipment"}


@frappe.whitelist()
def get_events(start, end, filters=None):
    """Get events for calendar view"""
    from frappe.desk.calendar import get_event_conditions
    
    conditions = get_event_conditions("Handover To Logistics", filters)
    
    events = frappe.db.sql("""
        SELECT 
            name,
            customer_name as subject,
            handover_date as start,
            handover_date as end,
            status,
            ekart_shipment_status,
            customer_name,
            order_no
        FROM `tabHandover To Logistics`
        WHERE handover_date BETWEEN %(start)s AND %(end)s
        {conditions}
    """.format(conditions=conditions), {
        "start": start,
        "end": end
    }, as_dict=True)
    
    # Color code events by status
    for event in events:
        event_status = event.ekart_shipment_status or event.status
        event.color = get_status_color(event_status)
    
    return events


def get_status_color(status):
    """Get color for status"""
    colors = {
        "Draft": "#d1d1d1",
        "Pending": "#ffa500",
        "Handed Over": "#6495ed",
        "In Transit": "#4169e1",
        "Delivered": "#32cd32",
        "Returned": "#dc143c",
        "Created": "#1e90ff",
        "Picked Up": "#00bfff",
        "Out for Delivery": "#ff8c00",
        "RTO": "#ff4500",
        "Exception": "#8b0000",
        "Cancelled": "#696969"
    }
    return colors.get(status, "#d1d1d1")


@frappe.whitelist()
def get_iawb_pending_dispatches():
    """Get all dispatches with IAWB Pending status for logistics handover

    This is used to populate the dispatch selection dropdown in Handover To Logistics
    """
    from impressio.material_out.doctype.dispatch_of_orders.dispatch_of_orders import get_iawb_pending_dispatches as get_dispatches
    return get_dispatches()


# File: /home/mdq/frappe-bench-v15/apps/inventre/inventre/material_out/doctype/handover_to_logistics/handover_to_logistics.py

@frappe.whitelist()
def create_handover_from_si(sales_invoice):
    """Create Handover To Logistics from Sales Invoice - ALL SHIPMENTS PREPAID"""
    try:
        # Get Sales Invoice
        si = frappe.get_doc("Sales Invoice", sales_invoice)
        
        # Check if already exists
        existing = frappe.db.exists("Handover To Logistics", {"sales_invoice": sales_invoice})
        if existing:
            frappe.throw(f"Handover already exists: {existing}")
        
        # VALIDATION: Ensure Sales Invoice is fully paid (since all shipments are prepaid)
        if si.outstanding_amount and si.outstanding_amount > 0:
            frappe.throw(
                f"Cannot create handover for Sales Invoice {sales_invoice} with outstanding amount {si.outstanding_amount}. "
                f"Please ensure the invoice is fully paid before creating handover."
            )
        
        # Get logistics settings
        settings = frappe.get_single("Logistics Settings")
        
        # Create new document
        handover = frappe.new_doc("Handover To Logistics")
        
        # === CRITICAL: SET STATUS TO DRAFT ===
        handover.status = "Draft"
        
        handover.logistics_partner = "" 
        
        # === BASIC INFORMATION ===
        handover.sales_invoice = sales_invoice
        handover.order_no = si.name
        handover.customer_name = si.customer_name or si.customer
        
        # Phone number - Check if field exists
        if hasattr(si, 'mobile_no') and si.mobile_no:
            handover.customer_phone = si.mobile_no
        elif hasattr(si, 'phone') and si.phone:
            handover.customer_phone = si.phone
        
        # === ADDRESS FROM SALES INVOICE ===
        address_name = si.shipping_address_name or si.customer_address
        if address_name:
            try:
                address = frappe.get_doc("Address", address_name)
                # Build address lines
                address_lines = []
                if address.address_line1:
                    address_lines.append(address.address_line1)
                if address.address_line2:
                    address_lines.append(address.address_line2)
                
                handover.address = "\n".join(address_lines)
                handover.pincode = address.pincode or settings.warehouse_pincode or "560001"
                handover.city = address.city or ""
                handover.state = address.state or ""
                handover.location = address.city or ""
            except:
                # Fallback if address not found
                handover.address = settings.warehouse_address or "Default Warehouse Address"
                handover.pincode = settings.warehouse_pincode or "560001"
        else:
            # Use settings warehouse address
            handover.address = settings.warehouse_address or "Default Warehouse Address"
            handover.pincode = settings.warehouse_pincode or "560001"
        
        # === PACKING MATERIALS FROM SALES INVOICE ITEMS ===
        for item in si.items:
            # Get item dimensions from item master if available
            item_master = frappe.db.get_value(
                "Item",
                item.item_code,
                ["pm_length", "pm_width", "pm_height", "pm_weight", "weight_per_unit"],
                as_dict=True
            ) or {}
            
            handover.append("packing_materials", {
                "item_code": item.item_code,
                "item_name": item.item_name,
                "qty": item.qty,
                "uom": item.uom,
                "pm_weight": float(
                    item_master.get("pm_weight") 
                    or item_master.get("weight_per_unit") 
                    or item.weight_per_unit 
                    or 0.5
                ),
                "pm_length": float(item_master.get("pm_length") or 0),
                "pm_width": float(item_master.get("pm_width") or 0),
                "pm_height": float(item_master.get("pm_height") or 0),
            })
        
        # === FINANCIAL VALUES ===
        handover.declared_value = si.grand_total or 0
        handover.invoice_value = si.grand_total or 0
        handover.invoice_number = si.name
        
        # FORCE ALL SHIPMENTS TO BE PREPAID - NO COD
        handover.payment_type = "Prepaid"
        handover.cod_amount = 0
        
        # Dates
        handover.handover_date = frappe.utils.today()
        
        # Save the document
        handover.insert(ignore_permissions=True)
        
        # Mark Sales Invoice
        frappe.db.set_value("Sales Invoice", sales_invoice, "handover_created", 1)
        frappe.db.commit()
        
        # Return success with redirect to EDIT mode
        return {
            "success": True,
            "handover": handover.name,
            "redirect": f"/app/handover-to-logistics/{handover.name}/edit",
            "message": f"Handover created: {handover.name}"
        }
        
    except Exception as e:
        import traceback
        frappe.log_error(
            title="Create Handover from SI Failed",
            message=f"SI: {sales_invoice}\nError: {str(e)}\n{traceback.format_exc()}"
        )
        frappe.throw(_("Failed to create handover: {0}").format(str(e)))



        
@frappe.whitelist()
def verify_handover_box(docname, barcode_or_qr):
    """API method to verify box during handover"""
    doc = frappe.get_doc("Handover To Logistics", docname)
    return doc.verify_box_by_scan(barcode_or_qr)


@frappe.whitelist()
def export_ekart_to_excel(docnames):
    """Export Ekart shipments to Excel in the required format"""
    import json
    from io import BytesIO
    import openpyxl
    from openpyxl.utils import get_column_letter
    from frappe.utils import now
    from frappe.utils.file_manager import save_file
    
    docnames = json.loads(docnames) if isinstance(docnames, str) else docnames
    
    # Get settings
    try:
        logistics_settings = frappe.get_single("Logistics Settings")
        pickup_facility = logistics_settings.get("ekart_pickup_facility") or "IEL_HYD_02"
        seller_name = logistics_settings.get("ekart_seller_name") or "Inventre Eduservices"
        gstin_number = logistics_settings.get("ekart_gstin") or "36AAMCP1199C1ZA"
    except:
        pickup_facility = "IEL_HYD_02"
        seller_name = "Inventre Eduservices"
        gstin_number = "36AAMCP1199C1ZA"
    
    data = []
    
    for docname in docnames:
        try:
            doc = frappe.get_doc("Handover To Logistics", docname)
            
            # Basic fields
            state = doc.get("state") or ""
            customer_phone = doc.get("customer_phone") or ""
            declared_value = float(doc.get("declared_value") or 0)
            length = float(doc.get("length") or 0)
            width = float(doc.get("width") or 0)
            height = float(doc.get("height") or 0)
            weight = float(doc.get("weight") or 0)
            
            # Address
            address = doc.get("address") or ""
            address_line_1 = ""
            address_line_2 = ""
            if address and "\n" in address:
                parts = address.split("\n", 1)
                address_line_1 = parts[0] or ""
                address_line_2 = parts[1] if len(parts) > 1 else ""
            else:
                address_line_1 = address
            
            # ========== GET SALES ORDER ==========
            order_id = ""
            invoice_number = ""
            
            if doc.get("sales_order"):
                order_id = doc.get("sales_order")
            
            if not order_id and doc.get("order_no"):
                order_no = doc.get("order_no")
                
                if order_no and order_no.startswith("DN-"):
                    try:
                        if frappe.db.exists("Delivery Note", order_no):
                            dn = frappe.get_doc("Delivery Note", order_no)
                            
                            if dn.get("items"):
                                for item in dn.items:
                                    so_value = item.get("custom_custom_against_sales_order")
                                    if so_value:
                                        order_id = so_value
                                        break
                            
                            if not order_id and dn.get("against_sales_order"):
                                order_id = dn.get("against_sales_order")
                            
                            if not order_id:
                                order_id = order_no
                    except Exception as e:
                        frappe.log_error(f"Error reading Delivery Note {order_no}: {str(e)}", "Ekart Export")
                
                elif order_no:
                    order_id = order_no
            
            if not order_id:
                order_id = docname
            
            invoice_number = order_id
            
            tracking_id = doc.get("logistics_tracking_number") or doc.get("tracking_number") or ""
            pincode = doc.get("pincode") or ""
            city = doc.get("city") or ""
            customer_name = doc.get("customer_name") or ""
            
            # Payment
            payment_type = doc.get("payment_type") or "Prepaid"
            cod_amount = float(doc.get("cod_amount") or 0)
            amount_to_collect = cod_amount if payment_type == "COD" else 0

            # ── Product info using custom_sub_items logic ──
            packing_items  = []
            parent_names   = {}   # deduped kit names
            main_items     = {}   # standalone item → qty
            total_quantity = 0
            product_title  = ""
            hsn_code       = ""
            category       = ""

            def _clean(name):
                if not name:
                    return ""
                if '$$' in name:
                    name = name.split('$$')[0]
                return name.lstrip(':').strip()

            # Resolve Sales Order
            so_id = doc.get("sales_order")
            if not so_id and doc.get("order_no"):
                if frappe.db.exists("Sales Order", doc.order_no):
                    so_id = doc.order_no
                elif frappe.db.exists("Delivery Note", doc.order_no):
                    try:
                        so_from_dn = frappe.db.get_value(
                            "Delivery Note Item",
                            {"parent": doc.order_no},
                            "custom_custom_against_sales_order"
                        )
                        if so_from_dn and frappe.db.exists("Sales Order", so_from_dn):
                            so_id = so_from_dn
                    except Exception:
                        pass

            if so_id:
                sub_items = []
                try:
                    if frappe.db.table_exists("Sale Order Sub Items"):
                        sub_items = frappe.db.sql("""
                            SELECT parent_item_code, item_code, qty
                            FROM `tabSale Order Sub Items`
                            WHERE parent = %s
                            ORDER BY idx ASC
                        """, so_id, as_dict=True)
                except Exception:
                    sub_items = []

                for row in sub_items:
                    p_code = (row.get("parent_item_code") or "").strip()
                    i_code = (row.get("item_code") or "").strip()
                    qty    = float(row.get("qty") or 1)
                    total_quantity += qty

                    label = _clean(p_code)
                    if not label:
                        continue

                    if p_code == i_code:
                        # standalone item
                        main_items[label] = main_items.get(label, 0) + qty
                    else:
                        # sub-item / kit
                        parent_names[label] = True

                    # HSN + category from item master (first hit wins)
                    if not hsn_code or not category:
                        try:
                            im = frappe.get_doc("Item", i_code)
                            if not hsn_code:
                                hsn_code = im.get("gst_hsn_code") or ""
                            if not category:
                                category = im.get("item_group") or ""
                        except Exception:
                            pass

            # Fallback to packing_materials
            if not parent_names and not main_items and doc.get("packing_materials"):
                for item in doc.packing_materials:
                    i_code    = item.get("item_code") or ""
                    item_name = (item.get("item_name") or i_code).strip()
                    qty       = float(item.get("qty") or 1)
                    total_quantity += qty
                    label = _clean(item_name)
                    if label:
                        main_items[label] = main_items.get(label, 0) + qty
                    if not hsn_code or not category:
                        try:
                            im = frappe.get_doc("Item", i_code)
                            if not hsn_code:
                                hsn_code = im.get("gst_hsn_code") or ""
                            if not category:
                                category = im.get("item_group") or ""
                        except Exception:
                            pass

            # Build display lists
            for name in parent_names:
                packing_items.append(name)
            for name, qty in main_items.items():
                packing_items.append(f"{name} x{int(qty)}")

            product_id    = ", ".join(packing_items) if packing_items else ""
            product_title = packing_items[0] if packing_items else customer_name
            if not category:
                category = "General"
            if not hsn_code:
                hsn_code = ""

            unit_price = declared_value / total_quantity if total_quantity > 0 else declared_value
            
            row = {
                "State": state,
                "Mobile Number": customer_phone,
                "Parent Tracking Id": "",
                "Declared Value": declared_value,
                "Return Label Line 1": pickup_facility,
                "Return Label Line 2": pickup_facility,
                "Length (CM)": length,
                "Breadth (CM)": width,
                "Height (CM)": height,
                "Shipment Weight (kg)": weight,
                "Product Id": product_id,
                "Category": category,
                "Product Title": product_title,
                "Quantity": int(total_quantity) if total_quantity > 0 else 1,
                "Unit Price": round(unit_price, 2),
                "CGST": 0,
                "SGST": 0,
                "IGST": 0,
                "HSN": hsn_code,
                "ERN": "",
                "Order Id": order_id,
                "Invoice Number": invoice_number,
                "Eway Bill Number": "",
                "Seller Name": seller_name,
                "GSTIN Number": gstin_number,
                "Tracking ID": tracking_id,
                "Amount To Collect": amount_to_collect,
                "Pickup Facility Name": pickup_facility,
                "Customer Name": customer_name,
                "Address Line 1": address_line_1,
                "Address Line 2": address_line_2,
                "Pincode": pincode,
                "City": city
            }
            
            data.append(row)
            
        except Exception as e:
            frappe.log_error(
                title="Error Processing Shipment for Export",
                message=f"Doc: {docname}\nError: {str(e)}"
            )
            continue
    
    if not data:
        frappe.throw(_("No data to export. Please check the selected shipments."))
    
    # Column order
    column_order = [
        "State", "Mobile Number", "Parent Tracking Id", "Declared Value",
        "Return Label Line 1", "Return Label Line 2", "Length (CM)", "Breadth (CM)",
        "Height (CM)", "Shipment Weight (kg)", "Product Id", "Category",
        "Product Title", "Quantity", "Unit Price", "CGST", "SGST", "IGST",
        "HSN", "ERN", "Order Id", "Invoice Number", "Eway Bill Number",
        "Seller Name", "GSTIN Number", "Tracking ID", "Amount To Collect",
        "Pickup Facility Name", "Customer Name", "Address Line 1", "Address Line 2",
        "Pincode", "City"
    ]
    
    # Create Excel file using openpyxl
    wb = openpyxl.Workbook()
    worksheet = wb.active
    worksheet.title = "Ekart Manifest"
    
    # Header row
    worksheet.append(column_order)
    
    # Data rows
    for row_dict in data:
        row_vals = [row_dict.get(col, "") for col in column_order]
        worksheet.append(row_vals)
        
    for column in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        for cell in column:
            try:
                if cell.value and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max(max_length + 2, 10), 50)
        worksheet.column_dimensions[column_letter].width = adjusted_width
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"Ekart_Manifest_{now()}.xlsx"
    file_doc = save_file(
        filename,
        output.getvalue(),
        "Handover To Logistics",
        docnames[0] if docnames else "Ekart_Export",
        decode=False,
        is_private=0
    )
    
    return {
        "file_url": file_doc.file_url,
        "filename": filename,
        "count": len(data)
    }